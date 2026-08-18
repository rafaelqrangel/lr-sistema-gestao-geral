# -*- coding: utf-8 -*-
"""Teste de ponta a ponta do modelo: preenche um cenario completo, recalcula no
LibreOffice e confere cada resultado contra o valor esperado calculado em Python.

Uso: python3 testar_modelo.py
"""
import shutil, subprocess, sys
from openpyxl import load_workbook

ORIG = "briefing_hppc_leite_de_rosas.xlsx"
COPIA = "/tmp/teste_modelo.xlsx"
RECALC = "/root/.claude/skills/synced/xlsx/scripts/recalc.py"

shutil.copy(ORIG, COPIA)
wb = load_workbook(COPIA)
b, m, f, e, c, v, s = (wb["1. Briefing"], wb["2. Mercado"], wb["3. Formulação"],
                       wb["4. Embalagem"], wb["5. Custos"], wb["6. Viabilidade"],
                       wb["7. Scorecard"])


def linha(ws, rotulo, col=1):
    for row in ws.iter_rows(min_col=col, max_col=col):
        val = row[0].value
        if isinstance(val, str) and val.strip() == rotulo:
            return row[0].row
    raise KeyError(f"{ws.title}: rotulo nao encontrado -> {rotulo}")


def põe(ws, rotulo, valor, col=2):
    ws.cell(row=linha(ws, rotulo), column=col).value = valor


# ============================================================ 1. BRIEFING
VOL1, MB_ALVO = 600_000, 0.35
põe(b, "Meta de volume — Ano 1 (un.)", VOL1)
põe(b, "Margem de contribuição-alvo (%)", MB_ALVO)
põe(b, "Capex máximo autorizado", 400_000)

# ============================================================= 2. MERCADO
POP, PEN, FREQ, PMED, COB, SHARE = 30_000_000, 0.40, 2, 20.0, 0.30, 0.02
for rot, val in [("População do público-alvo", POP), ("Penetração da categoria", PEN),
                 ("Frequência de compra", FREQ), ("Preço médio de gôndola da categoria", PMED),
                 ("Cobertura de distribuição alcançável", COB), ("Share-alvo dentro do SAM (Ano 1)", SHARE)]:
    põe(m, rot, val)
SOM = POP * PEN * FREQ * PMED * COB * SHARE
SOM_UN = SOM / PMED

# ========================================================== 3. FORMULAÇÃO
NVAR = 3
põe(f, "Nº de variações de fórmula a testar", NVAR)
BASICOS = [1264.0, 600.0, 1900.0, 1350.0]
CUSTO_BAS = NVAR * sum(BASICOS)
for rot, custo, prazo in [("Estabilidade acelerada", 3000.0, 45), ("Teste de uso em casa (HUT)", 8000.0, 60)]:
    r = linha(f, rot)
    f.cell(row=r, column=3).value = custo
    f.cell(row=r, column=5).value = prazo
CUSTO_ESP = NVAR * 11000.0
CUSTO_TST = CUSTO_BAS + CUSTO_ESP
reg0 = linha(f, "Enquadramento do grau de risco (1 ou 2)")
for i, val in enumerate([1500, 2500, 4000]):
    f.cell(row=reg0 + i, column=6).value = val
CUSTO_REG = 8000

# =========================================================== 4. EMBALAGEM
emb0 = linha(e, "Embalagem primária (frasco/pote/bisnaga)")
MOQS = [50000, 30000, 0, 20000, 0, 5000, 0, 0]
CAPEX = [45000, 15000, 0, 0, 0, 0, 0, 0]
for i, (mq, cp) in enumerate(zip(MOQS, CAPEX)):
    e.cell(row=emb0 + i, column=5).value = mq
    e.cell(row=emb0 + i, column=8).value = cp
EMB_CAPEX = sum(CAPEX)

# ============================================================== 5. CUSTOS
GRAM, OVER, PERDA_GR = 200.0, 0.01, 0.02
LOTE, VEL, HSETUP = 20_000, 2_000, 4.0
GRAM_EF = GRAM * (1 + OVER)
HPROD = LOTE / VEL
HTOT = HPROD + HSETUP
põe(c, "Gramatura declarada no rótulo", GRAM)
põe(c, "Sobre-enchimento médio", OVER)
põe(c, "Perda de granel no processo e no envase", PERDA_GR)
põe(c, "Tamanho do lote de produção", LOTE)
põe(c, "Velocidade da linha", VEL)
põe(c, "Horas de setup, limpeza e troca de lote", HSETUP)

# --- grupo 1: materias-primas (nome, %, preco bruto, frete, imp recuperaveis, perda)
MPS = [("Água purificada", 0.70, 0.02, 0.00, 0.00, 0.00),
       ("Ativo claim-driver", 0.05, 60.00, 1.50, 0.20, 0.02),
       ("Base emulsionante", 0.22, 4.00, 0.20, 0.12, 0.00),
       ("Fragrância", 0.03, 90.00, 0.00, 0.00, 0.01)]
mp0 = linha(c, "[EXEMPLO — apagar] Água purificada")
CUSTO_MP = 0.0
for i, (nome, pct, preco, frete, rec, perda) in enumerate(MPS):
    r = mp0 + i
    c.cell(row=r, column=1).value = nome
    c.cell(row=r, column=3).value = pct
    c.cell(row=r, column=5).value = preco
    c.cell(row=r, column=6).value = frete
    c.cell(row=r, column=7).value = rec
    c.cell(row=r, column=9).value = perda
    liquido = (preco + frete) * (1 - rec)
    CUSTO_MP += liquido * pct * GRAM_EF / 1000 / (1 - PERDA_GR) / (1 - perda)

# --- grupo 2: embalagem (qtd, preco, frete, imp recuperaveis, refugo)
EMBS = [("Embalagem primária (frasco, pote, bisnaga, lata)", 1, 0.80, 0.03, 0.12, 0.01),
        ("Tampa, válvula, pump ou dosador", 1, 0.25, 0.01, 0.12, 0.01),
        ("Rótulo, sleeve ou decoração", 1, 0.10, 0.00, 0.00, 0.02),
        ("Caixa de embarque", 1 / 12, 1.44, 0.00, 0.09, 0.00)]
CUSTO_EMB = 0.0
for nome, qtd, preco, frete, rec, refugo in EMBS:
    r = linha(c, nome)
    c.cell(row=r, column=3).value = qtd
    c.cell(row=r, column=5).value = preco
    c.cell(row=r, column=6).value = frete
    c.cell(row=r, column=7).value = rec
    c.cell(row=r, column=9).value = refugo
    CUSTO_EMB += (preco + frete) * (1 - rec) * qtd / (1 - refugo)

# --- grupo 3: MOD (pessoas, salario, fator, horas/mes)
MODS = [("Operador de envase", 2, 2000.0, 1.80, 176),
        ("Operador de embalagem e encaixotamento", 1, 1800.0, 1.80, 176)]
CUSTO_MOD = 0.0
for nome, n, sal, fator, hmes in MODS:
    r = linha(c, nome)
    c.cell(row=r, column=2).value = n
    c.cell(row=r, column=3).value = sal
    c.cell(row=r, column=4).value = fator
    c.cell(row=r, column=6).value = hmes
    CUSTO_MOD += (n * sal * fator) / hmes * HTOT / LOTE
for nome in ["Operador de fabricação do granel", "Auxiliar de linha", "Líder de linha"]:
    c.cell(row=linha(c, nome), column=2).value = 0

# --- grupo 4: GGF
GGFS = [("Depreciação de máquinas, moldes e instalações", 18000),
        ("Manutenção preventiva e corretiva", 7000),
        ("Energia elétrica da produção", 9000),
        ("Supervisão de produção e PCP", 11000),
        ("Limpeza, sanitização e produtos de higienização", 5000)]
for nome, val in GGFS:
    c.cell(row=linha(c, nome), column=3).value = val
GGF_MES = sum(v for _, v in GGFS)
GGF_HORAS, GGF_VAR = 500, 0.01
põe(c, "Horas produtivas da fábrica no mês", GGF_HORAS)
põe(c, "GGF variável medido por unidade", GGF_VAR)
CUSTO_GGF = GGF_MES / GGF_HORAS * HTOT / LOTE + GGF_VAR

# --- grupo 5: outros
QC_LOTE, AMOSTRAS, FERR_VAL, FERR_UN = 800.0, 10, 60_000.0, 500_000
ROY, OBSOL, TX_ANO, DIAS_EST = 0.02, 0.015, 0.18, 45
c.cell(row=linha(c, "Análises de qualidade por lote"), column=3).value = QC_LOTE
c.cell(row=linha(c, "Amostras de retenção e contraprova"), column=3).value = AMOSTRAS
rf = linha(c, "Amortização de ferramental e moldes")
c.cell(row=rf, column=3).value = FERR_VAL
c.cell(row=rf, column=5).value = FERR_UN
c.cell(row=linha(c, "Royalties e licenciamento"), column=3).value = ROY
c.cell(row=linha(c, "Perda por obsolescência e validade"), column=3).value = OBSOL
rg = linha(c, "Custo financeiro do estoque")
c.cell(row=rg, column=3).value = TX_ANO
c.cell(row=rg, column=5).value = DIAS_EST

# --- precificacao
PTAB, AIPI, AICMS, APIS, ACOF = 12.00, 0.05, 0.18, 0.0165, 0.076
AMVA, AICMSST, ADESC = 0.60, 0.18, 0.08
ACOM, AFRETE, AVERBA = 0.05, 0.03, 0.02
PRAT, MKVAR = 24.90, 0.40
for rot, val in [("Preço de tabela (sem IPI e sem ST)", PTAB), ("Alíquota de IPI", AIPI),
                 ("Alíquota de ICMS próprio", AICMS), ("Alíquota de PIS", APIS),
                 ("Alíquota de COFINS", ACOF), ("MVA / IVA-ST", AMVA),
                 ("Alíquota de ICMS no destino (para a ST)", AICMSST),
                 ("Descontos, bonificações e devoluções", ADESC),
                 ("Comissão de representantes", ACOM),
                 ("Frete de saída e armazenagem", AFRETE),
                 ("Verba de trade e ações de canal", AVERBA),
                 ("Preço de prateleira alvo (com impostos)", PRAT),
                 ("Markup do varejo", MKVAR),
                 ("Custos fixos incrementais mensais", 40_000),
                 ("Investimento de marketing mensal", 25_000)]:
    põe(c, rot, val)

MATERIAL = CUSTO_MP + CUSTO_EMB
CUSTO_OUT = (QC_LOTE / LOTE + AMOSTRAS / LOTE * MATERIAL + FERR_VAL / FERR_UN
             + ROY * PTAB + OBSOL * MATERIAL + MATERIAL * TX_ANO * DIAS_EST / 360)
CPV = CUSTO_MP + CUSTO_EMB + CUSTO_MOD + CUSTO_GGF + CUSTO_OUT

VIPI = PTAB * AIPI
BASEST = (PTAB + VIPI) * (1 + AMVA)
VST = max(0.0, BASEST * AICMSST - PTAB * AICMS)
NOTA = PTAB + VIPI + VST
RL = PTAB * (1 - AICMS - APIS - ACOF - ADESC)
MB = RL - CPV
MBP = MB / RL
DVAR = PTAB * (ACOM + AFRETE + AVERBA)
MC = MB - DVAR
MCP = MC / RL
AQUIS = PRAT / (1 + MKVAR)
FATOR = (1 + AIPI) * (1 + (1 + AMVA) * AICMSST) - AICMS
PTAB_IMP = AQUIS / FATOR
DEN = ((1 - AICMS - APIS - ACOF - ADESC - ACOM - AFRETE - AVERBA)
       - MB_ALVO * (1 - AICMS - APIS - ACOF - ADESC))
PTAB_MIN = CPV / DEN
PRAT_MIN = PTAB_MIN * FATOR * (1 + MKVAR)
BEUN = (40_000 + 25_000) / MC

# =========================================================== 6 e 7
G2, G3, PMKT, PCOM, TAXA = 0.25, 0.15, 0.08, 0.06, 0.18
põe(v, "Outros investimentos no Ano 0", 50_000)
NOTAS = [4, 4, 3, 5, 4, 4, 5, 3, 4, 3]
PESOS = [0.15, 0.15, 0.15, 0.10, 0.10, 0.08, 0.07, 0.10, 0.05, 0.05]
sc0 = linha(s, "Aderência à estratégia e à marca")
for i, n in enumerate(NOTAS):
    s.cell(row=sc0 + i, column=3).value = n
SCORE = sum(p * n for p, n in zip(PESOS, NOTAS))

EMB_1LOTE = max(MOQS) * CUSTO_EMB
INV = EMB_CAPEX + CUSTO_REG + CUSTO_TST + EMB_1LOTE + 50_000
VOL2, VOL3 = VOL1 * (1 + G2), VOL1 * (1 + G2) * (1 + G3)
RES = [(vol * RL) - (vol * CPV) - (vol * RL * PMKT) - (vol * RL * PCOM) for vol in (VOL1, VOL2, VOL3)]
VPL = -INV + sum(x / (1 + TAXA) ** (i + 1) for i, x in enumerate(RES))

wb.save(COPIA)
out = subprocess.run([sys.executable, RECALC, COPIA, "480"], capture_output=True, text=True)
print(out.stdout.strip())
if '"error"' in out.stdout:
    sys.exit(1)

# ============================================================ CONFERÊNCIA
wbv = load_workbook(COPIA, data_only=True)
b2, m2, f2, c2, v2, s2 = (wbv["1. Briefing"], wbv["2. Mercado"], wbv["3. Formulação"],
                          wbv["5. Custos"], wbv["6. Viabilidade"], wbv["7. Scorecard"])


def val(ws, rotulo, col=2):
    return ws.cell(row=linha(ws, rotulo), column=col).value


checks = [
    ("Gramatura efetiva envasada",   val(c2, "Gramatura efetiva envasada"), GRAM_EF),
    ("Horas totais do lote",         val(c2, "Horas totais ocupadas pelo lote"), HTOT),
    ("Grupo 1 · MP da fórmula",      val(c2, "Subtotal · matérias-primas da formulação", 10), CUSTO_MP),
    ("Grupo 2 · Embalagem",          val(c2, "Subtotal · embalagem", 10), CUSTO_EMB),
    ("Grupo 3 · MOD",                val(c2, "Subtotal · mão de obra direta", 10), CUSTO_MOD),
    ("Grupo 4 · GGF",                val(c2, "Total de GGF por unidade"), CUSTO_GGF),
    ("Grupo 5 · Outros",             val(c2, "Subtotal · outros custos de produção", 10), CUSTO_OUT),
    ("CPV unitário",                 val(c2, "CUSTO INDUSTRIAL TOTAL (CPV unitário)"), CPV),
    ("Custo do lote",                val(c2, "CUSTO INDUSTRIAL TOTAL (CPV unitário)", 4), CPV * LOTE),
    ("Valor do IPI",                 val(c2, "Valor do IPI"), VIPI),
    ("Base de cálculo da ST",        val(c2, "Base de cálculo da ST"), BASEST),
    ("ICMS-ST retido",               val(c2, "Valor do ICMS-ST retido"), VST),
    ("Valor total na nota",          val(c2, "VALOR TOTAL NA NOTA (o que o cliente paga)"), NOTA),
    ("Receita líquida",              val(c2, "RECEITA LÍQUIDA"), RL),
    ("Margem bruta (R$)",            val(c2, "MARGEM BRUTA"), MB),
    ("Margem bruta (%)",             val(c2, "Margem bruta (%)"), MBP),
    ("Margem de contribuição (R$)",  val(c2, "MARGEM DE CONTRIBUIÇÃO"), MC),
    ("Margem de contribuição (%)",   val(c2, "Margem de contribuição (%)"), MCP),
    ("Custo de aquisição do varejo", val(c2, "Custo de aquisição do varejo"), AQUIS),
    ("Fator nota → tabela",          val(c2, "Fator de conversão nota → tabela"), FATOR),
    ("Preço de tabela implícito",    val(c2, "Preço de tabela implícito no preço de prateleira"), PTAB_IMP),
    ("Preço de tabela mínimo",       val(c2, "Preço de tabela mínimo para a margem-alvo"), PTAB_MIN),
    ("Preço de prateleira mínimo",   val(c2, "Preço de prateleira mínimo"), PRAT_MIN),
    ("Break-even mensal",            val(c2, "Volume mensal de equilíbrio"), BEUN),
    ("SOM (R$)",                     val(m2, "SOM — receita capturável no Ano 1"), SOM),
    ("SOM (un.)",                    val(m2, "SOM em unidades"), SOM_UN),
    ("Plano de testes — total",      val(f2, "Custo total do plano de testes"), CUSTO_TST),
    ("Embalagem: 1º lote (aba 4)",   v2.cell(row=linha(v2, "Capital de giro — primeiro lote de embalagem"), column=2).value, EMB_1LOTE),
    ("Investimento total",           v2.cell(row=linha(v2, "► Investimento total"), column=2).value, INV),
    ("Resultado contributivo A1",    v2.cell(row=linha(v2, "► Resultado contributivo do projeto"), column=3).value, RES[0]),
    ("VPL",                          val(v2, "VPL (Valor Presente Líquido)"), VPL),
    ("Score ponderado",              s2.cell(row=linha(s2, "SCORE PONDERADO (1 a 5)"), column=4).value, SCORE),
    ("Painel: CPV (aba 1)",          val(b2, "Custo industrial unitário (CPV)"), CPV),
    ("Painel: nota (aba 1)",         val(b2, "Valor total na nota (com IPI e ST)"), NOTA),
    ("Painel: margem bruta (aba 1)", val(b2, "Margem bruta (%)"), MBP),
]

falhas = 0
print(f"\n{'indicador':32}{'planilha':>18}{'esperado':>18}   ok")
for nome, obtido, esperado in checks:
    ok = obtido is not None and abs(float(obtido) - float(esperado)) <= max(0.005, abs(esperado) * 1e-6)
    falhas += 0 if ok else 1
    print(f"{nome:32}{float(obtido) if obtido is not None else float('nan'):18,.4f}"
          f"{float(esperado):18,.4f}   {'OK' if ok else 'FALHOU'}")

print()
for nome, valor in [
    ("Fechamento do BOM",          c2.cell(row=linha(c2, "Validação do fechamento da fórmula"), column=3).value),
    ("Veredito de precificação",   val(c2, "VEREDITO DE PRECIFICAÇÃO")),
    ("Desvio do preço de tabela",  val(c2, "Desvio do preço de tabela praticado")),
    ("Participação MP no CPV",     val(c2, "Grupo 1 · Matérias-primas da formulação", 3)),
    ("Participação embalagem",     val(c2, "Grupo 2 · Embalagem", 3)),
    ("Participação MOD",           val(c2, "Grupo 3 · Mão de obra direta", 3)),
    ("Participação GGF",           val(c2, "Grupo 4 · Gastos gerais de fabricação", 3)),
    ("Classificação (aba 7)",      s2.cell(row=linha(s2, "CLASSIFICAÇÃO"), column=2).value),
    ("Capex x autorizado",         val(v2, "Investimento × capex autorizado")),
    ("Pendência de prazo",         val(f2, "Pendência de prazo")),
]:
    print(f"{nome:32} -> {valor}")

print(f"\n{len(checks)-falhas}/{len(checks)} verificacoes numericas OK")
sys.exit(1 if falhas else 0)
